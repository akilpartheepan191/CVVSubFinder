def substitutioner(absed,day):
  r_dict={}
  #definitions
  import openpyxl as xl
  import random
  from catter import catcon,categorise as catty
  def get_sheet_names(file_path):
    wb = xl.load_workbook(file_path)
    sheet_names = wb.sheetnames
    return sheet_names
  
  rel_teachers = {}
  
  def get_sub(day, period, category,teachers={}):
    if teachers=={}:
      file_name = 'time.xlsx'
      file = xl.load_workbook(file_name)
      days = {
        'mon': 'A',
        'tue': 'B',
        'wed': 'C',
        'thu': 'D',
        'fri': 'E',
        'sat': 'F',
      }
      for i in get_sheet_names(file_name):
        sheet = file[i]
        teachers[i] = []
        for y in range(8):
          p = y + 2
          teachers[i].append(sheet[f'{days[day]}{p}'].value)
        teachers[i].append(sheet['G2'].value)
      for i in list(teachers.keys()):
        sheet = file[i]
        no_of_pe = -1
        for pe in teachers[i]:
          if not None == pe:
            no_of_pe += 1
        teachers[i].append(no_of_pe)
        teachers[i].append(sheet['G3'].value)
      for i in list(teachers.keys()):
        if teachers[i][10]==1:
          del teachers[i]
      def remove_keys_from_dict(keys_list, dictionary):
        for key in keys_list:
            dictionary.pop(key, None)
      remove_keys_from_dict(absed,teachers)
    
    cat = category
    pe=period
    pe = pe - 1
    free = []
    for i in list(teachers.keys()):
      listee = teachers[i]
      if listee[pe] == None and catcon(cat,listee[8]):
        #print(listee[8])
        free.append(i)
  
    creds = {}
    for i in free:
      creds[i] = (teachers[i][9])
    def sort_dict_by_values(dictionary):
      sorted_dict = dict(sorted(dictionary.items(), key=lambda x: x[1]))
      return sorted_dict
  
    my_dict = creds
    sorted_dict = sort_dict_by_values(my_dict)
    new_creds = {}
    
    for key, value in sorted_dict.items():
      new_creds[key] = value
    if new_creds == {}:
      
      return 'noError'
    teachers_dict = new_creds 
    min_free_periods = min(teachers_dict.values())
    teachers_list = [teacher for teacher, free_periods in teachers_dict.items() if free_periods == min_free_periods]
    
    choosen_one = random.choice(teachers_list)
    teachers[choosen_one][9] = teachers[choosen_one][9] + 1
    return [teachers,choosen_one]
  
  def get_abs(day,absentees):
    teachers={}
    file_name = 'time.xlsx'
    file = xl.load_workbook(file_name)
    days = {
      'mon': 'A',
      'tue': 'B',
      'wed': 'C',
      'thu': 'D',
      'fri': 'E',
      'sat': 'F',
    }
    for i in get_sheet_names(file_name):
      sheet = file[i]
      teachers[i] = []
      for y in range(8):
        p = y + 2
        teachers[i].append(sheet[f'{days[day]}{p}'].value)
      teachers[i].append(sheet['G2'].value)
    for i in list(teachers.keys()):
      no_of_pe = -1
      for pe in teachers[i]:
        if not None == pe:
          no_of_pe += 1
      teachers[i].append(no_of_pe)
    absed={}
    for i in absentees:
      absed[i]=teachers[i]
    return absed
  def mapDict(theList,dict):
    newL=theList
    for i in list(dict.keys()):
      newL[i] = dict[i]
    return newL
  abs_data = get_abs(day,absed)
  
  for teacher in list(abs_data.keys()):
    
    r_dict[teacher]={}
    for period in range(0,8):
      teach_cat=abs_data[teacher][8]
      classs=abs_data[teacher][period]
      
      
      if not classs==None:
        rel_selection=get_sub(day,period+1,catty(classs),rel_teachers)
        if not rel_selection=="noError":
          selected=rel_selection[1]
          rel_teachers=rel_selection[0]
          r_dict[teacher][period+1]=[classs,selected]
          rel_teachers[selected][period]=classs
        else:
          r_dict[teacher][period+1]='No teacher is available!'
               
  return r_dict
  